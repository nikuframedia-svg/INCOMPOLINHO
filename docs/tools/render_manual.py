"""Render the Portuguese system manual to PDF and DOCX from one Markdown source.

Run in a documentation environment with python-docx, reportlab, markdown-it-py,
Pillow and PyMuPDF installed. These are not application runtime dependencies.
"""

from __future__ import annotations

import html
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from markdown_it import MarkdownIt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    Image,
    KeepTogether,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)
from reportlab.platypus.tableofcontents import TableOfContents


DOCS = Path(__file__).resolve().parents[1]
SOURCE = DOCS / "manual-tecnico-funcional-incompol.md"
OUT = DOCS / "entregaveis"
BASENAME = "INCOMPOL_Manual_Tecnico_Funcional_v1.0"
TITLE = "Manual técnico e funcional"
SUBTITLE = "Sistema de planeamento de produção"
DATE = "16 de setembro de 2026"
MAX_PAGES = 20
INK = "173B3B"
ACCENT = "007E78"
MUTED = "5B6971"
PALE = "EEF5F4"


def hex_color(value):
    return colors.HexColor("#" + value.lstrip("#"))


def parse_blocks():
    tokens = MarkdownIt("commonmark").enable("table").parse(SOURCE.read_text())
    blocks = []
    started = False
    list_stack = []
    i = 0
    while i < len(tokens):
        token = tokens[i]
        if token.type == "heading_open":
            content = tokens[i + 1]
            if content.content.startswith("1. Objetivo"):
                started = True
            if started:
                blocks.append({"kind": "heading", "level": int(token.tag[1:]) - 1,
                               "text": content.content, "inline": content.children})
            i += 3
            continue
        if not started:
            i += 1
            continue
        if token.type in {"bullet_list_open", "ordered_list_open"}:
            list_stack.append({"ordered": token.type == "ordered_list_open",
                               "number": int(token.attrGet("start") or 1)})
        elif token.type in {"bullet_list_close", "ordered_list_close"}:
            list_stack.pop()
        elif token.type == "paragraph_open":
            inline = tokens[i + 1]
            images = [child for child in inline.children or [] if child.type == "image"]
            if images:
                for image in images:
                    blocks.append({"kind": "image", "path": DOCS / image.attrGet("src"),
                                   "text": image.content})
            else:
                prefix = ""
                if list_stack:
                    entry = list_stack[-1]
                    prefix = f"{entry['number']}. " if entry["ordered"] else "• "
                    entry["number"] += 1
                blocks.append({"kind": "paragraph", "inline": inline.children,
                               "text": inline.content, "prefix": prefix})
            i += 3
            continue
        elif token.type == "fence":
            blocks.append({"kind": "code", "text": token.content.rstrip()})
        elif token.type == "table_open":
            rows = []
            row = []
            i += 1
            while tokens[i].type != "table_close":
                current = tokens[i]
                if current.type == "tr_open":
                    row = []
                elif current.type == "inline":
                    row.append(current.children or [])
                elif current.type == "tr_close":
                    rows.append(row)
                i += 1
            blocks.append({"kind": "table", "rows": rows})
        i += 1
    return blocks


def plain(inline):
    return "".join(token.content if token.type in {"text", "code_inline"} else
                   " " if token.type in {"softbreak", "hardbreak"} else ""
                   for token in inline or [])


def pdf_inline(inline):
    result = []
    for token in inline or []:
        if token.type == "text":
            result.append(html.escape(token.content))
        elif token.type == "code_inline":
            result.append(f'<font name="ManualMono" size="8.5">{html.escape(token.content)}</font>')
        elif token.type == "strong_open":
            result.append("<b>")
        elif token.type == "strong_close":
            result.append("</b>")
        elif token.type == "em_open":
            result.append("<i>")
        elif token.type == "em_close":
            result.append("</i>")
        elif token.type in {"softbreak", "hardbreak"}:
            result.append("<br/>" if token.type == "hardbreak" else " ")
    return "".join(result)


def column_fractions(rows):
    count = len(rows[0])
    first_header = plain(rows[0][0])
    if count == 3 and first_header in {"ID", "Código"}:
        return [.09, .33, .58]
    if count == 2 and first_header == "Termo":
        return [.27, .73]
    return {2: [.36, .64], 3: [.23, .28, .49], 4: [.22, .24, .23, .31],
            5: [.19, .19, .20, .22, .20]}.get(count, [1 / count] * count)


def register_fonts():
    root = Path("/usr/share/fonts/truetype/liberation")
    for name, filename in [
        ("Manual", "LiberationSans-Regular.ttf"),
        ("ManualBold", "LiberationSans-Bold.ttf"),
        ("ManualItalic", "LiberationSans-Italic.ttf"),
        ("ManualBoldItalic", "LiberationSans-BoldItalic.ttf"),
        ("ManualMono", "LiberationMono-Regular.ttf"),
    ]:
        pdfmetrics.registerFont(TTFont(name, str(root / filename)))
    pdfmetrics.registerFontFamily("Manual", normal="Manual", bold="ManualBold",
                                  italic="ManualItalic", boldItalic="ManualBoldItalic")


class ManualPDF(BaseDocTemplate):
    def __init__(self, path):
        super().__init__(str(path), pagesize=A4, leftMargin=19 * mm, rightMargin=19 * mm,
                         topMargin=21 * mm, bottomMargin=20 * mm,
                         title=f"INCOMPOL | {TITLE}", author="INCOMPOL / ProdPlan ONE")
        frame = Frame(self.leftMargin, self.bottomMargin, self.width, self.height,
                      leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
        self.addPageTemplates(PageTemplate(id="manual", frames=[frame], onPage=self.chrome))
        self.heading_pages = {}

    def chrome(self, canvas, doc):
        canvas.saveState()
        if doc.page == 1:
            canvas.setFillColor(hex_color(INK))
            canvas.rect(0, A4[1] - 66 * mm, A4[0], 66 * mm, fill=1, stroke=0)
            canvas.setFillColor(hex_color(ACCENT))
            canvas.rect(19 * mm, A4[1] - 70 * mm, 24 * mm, 2 * mm, fill=1, stroke=0)
        else:
            canvas.setFont("ManualBold", 8)
            canvas.setFillColor(hex_color(INK))
            canvas.drawString(19 * mm, A4[1] - 12 * mm, "INCOMPOL  /  PRODPLAN ONE")
            canvas.setFont("Manual", 8)
            canvas.setFillColor(hex_color(MUTED))
            canvas.drawRightString(A4[0] - 19 * mm, A4[1] - 12 * mm, TITLE)
            canvas.setStrokeColor(hex_color("D8E3E0"))
            canvas.line(19 * mm, A4[1] - 15 * mm, A4[0] - 19 * mm, A4[1] - 15 * mm)
        canvas.setFont("Manual", 8)
        canvas.setFillColor(hex_color(MUTED))
        canvas.drawString(19 * mm, 11 * mm, "INCOMPOL-MTF-001  |  Edição 1.0  |  16-09-2026")
        canvas.drawRightString(A4[0] - 19 * mm, 11 * mm, str(doc.page))
        canvas.restoreState()

    def afterFlowable(self, flowable):
        if isinstance(flowable, Paragraph) and hasattr(flowable, "manual_heading"):
            level, title, key = flowable.manual_heading
            self.canv.bookmarkPage(key)
            self.canv.addOutlineEntry(title, key, level=level, closed=level > 0)
            self.notify("TOCEntry", (level, title, self.page, key))
            self.heading_pages[title] = self.page


def make_pdf(blocks):
    register_fonts()
    body = ParagraphStyle("Body", fontName="Manual", fontSize=10.2, leading=14.2,
                          textColor=hex_color("253238"), spaceAfter=7, splitLongWords=True,
                          allowWidows=0, allowOrphans=0)
    heading1 = ParagraphStyle("Major", parent=body, fontName="ManualBold", fontSize=20,
                              leading=25, spaceBefore=0, spaceAfter=14, keepWithNext=True,
                              textColor=hex_color(INK))
    heading2 = ParagraphStyle("Minor", parent=body, fontName="ManualBold", fontSize=12.5,
                              leading=16.5, spaceBefore=13, spaceAfter=7, keepWithNext=True,
                              textColor=hex_color(ACCENT))
    cell = ParagraphStyle("Cell", parent=body, fontSize=8.8, leading=12, spaceAfter=0)
    header = ParagraphStyle("CellHeader", parent=cell, fontName="ManualBold", textColor=colors.white)
    caption = ParagraphStyle("Caption", parent=body, fontSize=8.7, leading=12,
                             textColor=hex_color(MUTED), spaceBefore=5, spaceAfter=12)
    code = ParagraphStyle("Formula", parent=body, fontName="ManualMono", fontSize=8.7,
                          leading=12.8, backColor=hex_color(PALE), borderPadding=9,
                          spaceBefore=5, spaceAfter=12)
    toc = TableOfContents()
    toc.levelStyles = [
        ParagraphStyle("TOC1", parent=body, fontName="ManualBold", fontSize=10.2, leading=14,
                       spaceBefore=7, spaceAfter=3, rightIndent=18),
        ParagraphStyle("TOC2", parent=body, fontSize=9, leading=12, leftIndent=12,
                       firstLineIndent=0, spaceAfter=2, rightIndent=18),
    ]
    cover_brand = ParagraphStyle("Brand", parent=body, fontName="ManualBold", fontSize=33,
                                 leading=39, textColor=colors.white)
    cover_small = ParagraphStyle("CoverSmall", parent=body, fontSize=11, leading=16,
                                 textColor=hex_color("B9DCD6"))
    cover_title = ParagraphStyle("CoverTitle", parent=heading1, fontSize=32, leading=38)
    story = [Spacer(1, 8 * mm), Paragraph("INCOMPOL", cover_brand),
             Paragraph("PRODPLAN ONE", cover_small), Spacer(1, 49 * mm),
             Paragraph("Manual técnico<br/>e funcional", cover_title),
             Paragraph(SUBTITLE, ParagraphStyle("Subtitle", parent=body, fontSize=17, leading=23)),
             Spacer(1, 14 * mm),
             Paragraph("Regras de negócio · Indicadores · Ecrãs · Funcionalidades", body),
             Paragraph("Operação · Otimização · Arquitetura · Limitações", body),
             Spacer(1, 32 * mm), Paragraph(f"<b>Edição 1.0</b><br/>{DATE}", body),
             Paragraph("Referência técnica e operacional baseada no código e na configuração do sistema.", caption),
             PageBreak(), Paragraph("Índice", heading1), toc, PageBreak()]
    first_heading = True
    heading_id = 0
    width = A4[0] - 38 * mm
    for block in blocks:
        kind = block["kind"]
        if kind == "heading":
            level = block["level"]
            if level == 1 and not first_heading:
                story.append(PageBreak())
            first_heading = False
            paragraph = Paragraph(html.escape(block["text"]), heading1 if level == 1 else heading2)
            paragraph.manual_heading = (min(level - 1, 1), block["text"], f"h{heading_id}")
            heading_id += 1
            story.append(paragraph)
        elif kind == "paragraph":
            content = html.escape(block["prefix"]) + pdf_inline(block["inline"])
            style = body
            if block["prefix"]:
                style = ParagraphStyle("List", parent=body, leftIndent=13, firstLineIndent=-13,
                                       spaceAfter=4)
            story.append(Paragraph(content, style))
        elif kind == "code":
            lines = [html.escape(line).replace("  ", "&#160;&#160;") for line in block["text"].splitlines()]
            story.append(Paragraph("<br/>".join(lines), code))
        elif kind == "table":
            fractions = column_fractions(block["rows"])
            padding = 4.5 if plain(block["rows"][0][0]) == "Termo" else 5
            data = [[Paragraph(pdf_inline(tokens), header if row_idx == 0 else cell)
                     for tokens in row] for row_idx, row in enumerate(block["rows"])]
            table = Table(data, colWidths=[width * f for f in fractions], repeatRows=1,
                          hAlign="LEFT", spaceBefore=4, spaceAfter=12)
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), hex_color(INK)),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, hex_color("F4F7F6")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), padding),
                ("BOTTOMPADDING", (0, 0), (-1, -1), padding),
                ("LINEBELOW", (0, 0), (-1, 0), .6, hex_color(ACCENT)),
                ("LINEBELOW", (0, 1), (-1, -1), .3, hex_color("D9E3E0")),
            ]))
            story.append(table)
        elif kind == "image":
            image = Image(str(block["path"]), width=width, height=width * 960 / 1440)
            group = [image, Paragraph(html.escape(block["text"]), caption)]
            if story and hasattr(story[-1], "manual_heading"):
                group.insert(0, story.pop())
            story.append(KeepTogether(group))
    pdf = ManualPDF(OUT / f"{BASENAME}.pdf")
    pdf.multiBuild(story)
    if pdf.page > MAX_PAGES:
        raise ValueError(f"Manual exceeds {MAX_PAGES} pages: {pdf.page}")
    print(f"PDF: {pdf.page} pages")


def docx_inline(paragraph, tokens):
    bold = italic = False
    for token in tokens or []:
        if token.type == "strong_open":
            bold = True
        elif token.type == "strong_close":
            bold = False
        elif token.type == "em_open":
            italic = True
        elif token.type == "em_close":
            italic = False
        elif token.type in {"text", "code_inline", "softbreak", "hardbreak"}:
            value = "\n" if token.type == "hardbreak" else " " if token.type == "softbreak" else token.content
            run = paragraph.add_run(value)
            run.bold, run.italic = bold, italic
            if token.type == "code_inline":
                run.font.name = "Liberation Mono"
                run.font.size = Pt(8.5)


def shade_cell(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    properties.append(shade)


def make_docx(blocks):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2), Cm(1.9)
    section.left_margin = section.right_margin = Cm(1.9)
    section.header_distance = section.footer_distance = Cm(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Liberation Sans"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("253238")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.13
    for level, size in [(1, 20), (2, 13), (3, 11)]:
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Liberation Sans"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(INK if level == 1 else ACCENT)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.space_before = Pt(14)
    doc.styles["Caption"].font.name = "Liberation Sans"
    doc.styles["Caption"].font.size = Pt(9)
    doc.styles["Caption"].font.color.rgb = RGBColor.from_string(MUTED)
    header = section.header.paragraphs[0]
    header.text = "INCOMPOL  /  PRODPLAN ONE                                        Manual técnico e funcional"
    header.style = doc.styles["Caption"]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("INCOMPOL-MTF-001  |  Edição 1.0  |  ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(35)
    run = p.add_run("INCOMPOL")
    run.bold = True
    run.font.size = Pt(38)
    run.font.color.rgb = RGBColor.from_string(INK)
    doc.add_paragraph("PRODPLAN ONE", "Subtitle")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(80)
    run = p.add_run("Manual técnico\ne funcional")
    run.bold = True
    run.font.size = Pt(32)
    run.font.color.rgb = RGBColor.from_string(INK)
    doc.add_paragraph(SUBTITLE, "Subtitle")
    doc.add_paragraph("Regras de negócio · Indicadores · Ecrãs · Funcionalidades")
    doc.add_paragraph("Operação · Otimização · Arquitetura · Limitações")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(70)
    p.add_run(f"Edição 1.0\n{DATE}").bold = True
    doc.add_paragraph("Referência técnica e operacional baseada no código e na configuração do sistema.", "Caption")
    doc.add_page_break()
    doc.add_paragraph("Índice navegável", "Title")
    headings = [block for block in blocks if block["kind"] == "heading"]
    for idx, block in enumerate(headings):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Cm(0.4 if block["level"] > 1 else 0)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("w:anchor"), f"manual_{idx}")
        run = OxmlElement("w:r")
        properties = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), INK if block["level"] == 1 else MUTED)
        properties.append(color)
        if block["level"] == 1:
            properties.append(OxmlElement("w:b"))
        run.append(properties)
        text = OxmlElement("w:t")
        text.text = block["text"]
        run.append(text)
        hyperlink.append(run)
        p._p.append(hyperlink)
    doc.add_page_break()
    heading_index = 0
    for block in blocks:
        kind = block["kind"]
        if kind == "heading":
            p = doc.add_heading(block["text"], level=min(3, block["level"]))
            p.paragraph_format.page_break_before = block["level"] == 1 and heading_index > 0
            start = OxmlElement("w:bookmarkStart")
            start.set(qn("w:id"), str(heading_index))
            start.set(qn("w:name"), f"manual_{heading_index}")
            end = OxmlElement("w:bookmarkEnd")
            end.set(qn("w:id"), str(heading_index))
            p._p.insert(0, start)
            p._p.append(end)
            heading_index += 1
        elif kind == "paragraph":
            p = doc.add_paragraph()
            if block["prefix"]:
                p.add_run(block["prefix"])
                p.paragraph_format.left_indent = Cm(.45)
                p.paragraph_format.first_line_indent = Cm(-.45)
            docx_inline(p, block["inline"])
        elif kind == "code":
            p = doc.add_paragraph()
            run = p.add_run(block["text"])
            run.font.name, run.font.size = "Liberation Mono", Pt(8.5)
            p.paragraph_format.space_before = Pt(7)
            p.paragraph_format.space_after = Pt(10)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), PALE)
            p._p.get_or_add_pPr().append(shade)
        elif kind == "table":
            rows = block["rows"]
            table = doc.add_table(rows=0, cols=len(rows[0]))
            table.autofit = False
            fractions = column_fractions(rows)
            for column, fraction in zip(table.columns, fractions):
                column.width = Cm(17.2 * fraction)
            for row_idx, values in enumerate(rows):
                row = table.add_row()
                cant_split = OxmlElement("w:cantSplit")
                row._tr.get_or_add_trPr().append(cant_split)
                if row_idx == 0:
                    row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
                for idx, (cell, tokens) in enumerate(zip(row.cells, values)):
                    cell.width = Cm(17.2 * fractions[idx])
                    shade_cell(cell, INK if row_idx == 0 else "F2F6F5" if row_idx % 2 == 0 else "FFFFFF")
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.line_spacing = 1.06
                    docx_inline(p, tokens)
                    for run in p.runs:
                        run.font.size = Pt(9)
                        if row_idx == 0:
                            run.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif kind == "image":
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(block["path"]), width=Cm(17.2))
            doc.add_paragraph(block["text"], "Caption")
    doc.core_properties.title = f"INCOMPOL | {TITLE}"
    doc.core_properties.subject = SUBTITLE
    doc.core_properties.author = "INCOMPOL / ProdPlan ONE"
    doc.core_properties.keywords = "INCOMPOL, ISOP, APS, OTD, OTD-D, planeamento, manual"
    doc.save(OUT / f"{BASENAME}.docx")
    print(f"DOCX: {len(headings)} indexed sections; {len(doc.tables)} tables")


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    blocks = parse_blocks()
    make_pdf(blocks)
    make_docx(blocks)
